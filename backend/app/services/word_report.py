from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.models.schemas import LoadMode, MetricsSnapshot, RunDetail, RunStatus


def _fmt_dt(value: datetime | None) -> str:
    if value is None:
        return "-"
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    local = value.astimezone()
    return local.strftime("%Y-%m-%d %H:%M:%S %Z")


def _fmt_ms(value: float | None) -> str:
    if value is None:
        return "-"
    return f"{value:.0f}"


def _fmt_pct(ratio: float) -> str:
    return f"{ratio * 100:.2f}%"


def _duration_seconds(detail: RunDetail) -> float | None:
    if detail.started_at and detail.finished_at:
        return max(0.0, (detail.finished_at - detail.started_at).total_seconds())
    return None


def _history_peaks(history: list[MetricsSnapshot]) -> dict[str, Any]:
    if not history:
        return {
            "peak_rps": 0.0,
            "peak_users": 0,
            "peak_p95": None,
            "peak_fail_ratio": 0.0,
            "avg_rps": 0.0,
        }
    peak_rps_point = max(history, key=lambda h: h.total_rps)
    peak_user_point = max(history, key=lambda h: h.user_count)
    with_p95 = [h for h in history if h.p95 is not None]
    peak_p95_point = max(with_p95, key=lambda h: h.p95 or 0.0) if with_p95 else None
    peak_fail_point = max(history, key=lambda h: h.fail_ratio)
    avg_rps = sum(h.total_rps for h in history) / len(history)
    return {
        "peak_rps": peak_rps_point.total_rps,
        "peak_rps_at": peak_rps_point.ts,
        "peak_users": peak_user_point.user_count,
        "peak_users_at": peak_user_point.ts,
        "peak_p95": peak_p95_point.p95 if peak_p95_point else None,
        "peak_p95_at": peak_p95_point.ts if peak_p95_point else None,
        "peak_fail_ratio": peak_fail_point.fail_ratio,
        "avg_rps": avg_rps,
    }


def _status_label(status: RunStatus) -> str:
    mapping = {
        RunStatus.finished: "正常完成",
        RunStatus.stopped: "人工停止",
        RunStatus.failed: "失败",
        RunStatus.running: "运行中",
        RunStatus.queued: "排队中",
        RunStatus.stopping: "停止中",
    }
    return mapping.get(status, status.value)


def _build_narrative(detail: RunDetail, peaks: dict[str, Any]) -> list[str]:
    m = detail.latest_metrics
    elapsed = _duration_seconds(detail)
    elapsed_text = f"{elapsed:.1f} 秒" if elapsed is not None else "未知"
    lines: list[str] = []

    if detail.load_mode == LoadMode.throughput and detail.target_rps:
        mode_text = (
            f"负载模式为恒定吞吐，目标约 {detail.target_rps:.2f} RPS，"
            f"并发用户池 {detail.users}（用于承载吞吐）"
        )
    else:
        mode_text = f"负载模式为恒定用户，并发用户 {detail.users}"
    lines.append(
        f"本次压测针对目标主机 {detail.host}，使用场景「{detail.scenario_name}」"
        f"（{detail.scenario_id}）。{mode_text}，启动速率 {detail.spawn_rate}/s、"
        f"Worker 数 {detail.workers}、计划时长 {detail.duration}。"
        f"实际运行状态为「{_status_label(detail.status)}」，有效时长约 {elapsed_text}。"
    )

    paths = "、".join(detail.paths) if detail.paths else "/"
    lines.append(
        f"压测路径包括：{paths}。"
        + (
            "已开启静态资源跟随请求。"
            if detail.follow_static_assets
            else "未跟随页面内静态资源。"
        )
    )

    if not m:
        lines.append("未能采集到有效统计指标，请检查 Locust 进程日志与目标可达性。")
        if detail.error:
            lines.append(f"错误信息：{detail.error}")
        return lines

    lines.append(
        f"累计完成请求 {m.total_requests} 次，失败 {m.total_failures} 次，"
        f"失败率 {_fmt_pct(m.fail_ratio)}。结束时瞬时 RPS 为 {m.total_rps:.2f}，"
        f"历史峰值 RPS 为 {peaks['peak_rps']:.2f}，全程平均 RPS 约 {peaks['avg_rps']:.2f}。"
        f"延迟分位：p50={_fmt_ms(m.p50)} ms，p95={_fmt_ms(m.p95)} ms，p99={_fmt_ms(m.p99)} ms。"
    )

    if m.fail_ratio <= 0.001 and (m.p95 is None or m.p95 < 500):
        lines.append(
            "从客户端观测看，目标在本次配置下保持较低失败率与可接受延迟，"
            "整体表现相对平稳。需注意：指标来自发压端统计，是否完全到达源站应用层"
            "仍应以目标侧 access log / APM 为准。"
        )
    elif m.fail_ratio > 0.01:
        lines.append(
            "失败率偏高，建议结合失败请求明细、目标站错误日志与发压机资源占用，"
            "区分是目标容量不足、限流/WAF，还是客户端超时或网络问题。"
        )
    elif m.p95 is not None and m.p95 >= 1000:
        lines.append(
            "p95 延迟偏高，说明在接近峰值负载时响应明显变慢；若发压机 CPU/网卡仍有余量，"
            "瓶颈更可能位于目标站或其前置 CDN/网关。"
        )
    else:
        lines.append(
            "整体指标处于中等波动区间。建议结合峰值时段与分接口统计进一步定位慢点。"
        )

    if detail.error:
        lines.append(f"运行过程记录到错误：{detail.error}")

    lines.append(
        "本报告由本地压测工作台自动生成，仅用于授权范围内的性能评估，"
        "请勿对未授权系统发起压力测试。"
    )
    return lines


def _set_run_font(document: Document) -> None:
    style = document.styles["Normal"]
    font = style.font
    font.name = "Microsoft YaHei"
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_kv_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for idx, (key, value) in enumerate(rows):
        table.rows[idx].cells[0].text = key
        table.rows[idx].cells[1].text = value
        for cell in table.rows[idx].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_word_report(detail: RunDetail, output_path: Path) -> Path:
    """Generate a Word (.docx) report for a finished/stopped/failed run."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    peaks = _history_peaks(detail.history)
    m = detail.latest_metrics
    elapsed = _duration_seconds(detail)

    document = Document()
    _set_run_font(document)
    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("压力测试报告")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = subtitle.add_run(f"Run ID: {detail.id}")
    sub.font.size = Pt(11)
    sub.font.name = "Consolas"

    _add_heading(document, "一、基本信息", level=1)
    _add_kv_table(
        document,
        [
            ("场景名称", detail.scenario_name),
            ("场景 ID", detail.scenario_id),
            ("目标 Host", detail.host),
            (
                "负载模式",
                "恒定吞吐 (throughput)"
                if detail.load_mode == LoadMode.throughput
                else "恒定用户 (users)",
            ),
            ("并发用户数", str(detail.users)),
            (
                "目标 RPS",
                f"{detail.target_rps:.2f}" if detail.target_rps is not None else "-",
            ),
            ("启动速率 (users/s)", str(detail.spawn_rate)),
            ("Worker 数", str(detail.workers)),
            ("计划时长", detail.duration),
            ("实际时长", f"{elapsed:.1f} s" if elapsed is not None else "-"),
            ("压测路径", "、".join(detail.paths) if detail.paths else "/"),
            ("跟随静态资源", "是" if detail.follow_static_assets else "否"),
            ("运行状态", f"{_status_label(detail.status)} ({detail.status.value})"),
            ("创建时间", _fmt_dt(detail.created_at)),
            ("开始时间", _fmt_dt(detail.started_at)),
            ("结束时间", _fmt_dt(detail.finished_at)),
            ("错误信息", detail.error or "无"),
        ],
    )

    _add_heading(document, "二、统计数据总览", level=1)
    _add_kv_table(
        document,
        [
            ("总请求数", str(m.total_requests if m else 0)),
            ("总失败数", str(m.total_failures if m else 0)),
            ("失败率", _fmt_pct(m.fail_ratio) if m else "0.00%"),
            ("结束时 RPS", f"{m.total_rps:.2f}" if m else "0.00"),
            ("峰值 RPS", f"{peaks['peak_rps']:.2f}"),
            ("全程平均 RPS", f"{peaks['avg_rps']:.2f}"),
            ("结束时活跃用户", str(m.user_count if m else 0)),
            ("峰值活跃用户", str(peaks["peak_users"])),
            ("p50 延迟 (ms)", _fmt_ms(m.p50 if m else None)),
            ("p95 延迟 (ms)", _fmt_ms(m.p95 if m else None)),
            ("p99 延迟 (ms)", _fmt_ms(m.p99 if m else None)),
            ("历史峰值 p95 (ms)", _fmt_ms(peaks["peak_p95"])),
            ("历史峰值失败率", _fmt_pct(float(peaks["peak_fail_ratio"]))),
            ("采样点数", str(len(detail.history))),
        ],
    )

    _add_heading(document, "三、详细情况描述", level=1)
    for paragraph in _build_narrative(detail, peaks):
        document.add_paragraph(paragraph)

    _add_heading(document, "四、关键时间点", level=1)
    bullets: list[str] = []
    if peaks.get("peak_rps_at"):
        bullets.append(
            f"峰值 RPS {peaks['peak_rps']:.2f} 出现于 {_fmt_dt(peaks['peak_rps_at'])}"
        )
    if peaks.get("peak_users_at"):
        bullets.append(
            f"峰值用户数 {peaks['peak_users']} 出现于 {_fmt_dt(peaks['peak_users_at'])}"
        )
    if peaks.get("peak_p95_at") and peaks.get("peak_p95") is not None:
        bullets.append(
            f"峰值 p95 {_fmt_ms(peaks['peak_p95'])} ms 出现于 {_fmt_dt(peaks['peak_p95_at'])}"
        )
    if bullets:
        for item in bullets:
            document.add_paragraph(item, style="List Bullet")
    else:
        document.add_paragraph("无可用时序采样点。")

    endpoints = m.endpoints if m else []
    _add_heading(document, "五、分接口明细", level=1)
    if not endpoints:
        document.add_paragraph("无接口级统计数据。")
    else:
        table = document.add_table(rows=1 + len(endpoints), cols=9)
        table.style = "Table Grid"
        headers = [
            "接口",
            "方法",
            "请求数",
            "失败数",
            "瞬时 RPS",
            "平均耗时(ms)",
            "p50",
            "p95",
            "p99",
        ]
        for i, text in enumerate(headers):
            table.rows[0].cells[i].text = text
        for row_idx, ep in enumerate(endpoints, start=1):
            values = [
                ep.name,
                ep.method,
                str(ep.num_requests),
                str(ep.num_failures),
                f"{ep.current_rps:.1f}",
                f"{ep.avg_response_time:.1f}",
                _fmt_ms(ep.p50),
                _fmt_ms(ep.p95),
                _fmt_ms(ep.p99),
            ]
            for col, value in enumerate(values):
                table.rows[row_idx].cells[col].text = value
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Microsoft YaHei"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    _add_heading(document, "六、说明与免责", level=1)
    document.add_paragraph(
        "1. RPS、延迟、失败率均由 Locust 发压端统计，表示客户端观测到的完成请求情况。"
    )
    document.add_paragraph(
        "2. 若目标前有 CDN、WAF 或负载均衡，边缘节点处理量可能与源站应用层处理量不一致。"
    )
    document.add_paragraph(
        "3. 请仅对已获授权的系统进行压测；本工具不构成对目标可用性或容量的法律保证。"
    )

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run(f"生成时间：{_fmt_dt(datetime.now(timezone.utc))}")
    fr.font.size = Pt(9)
    fr.font.name = "Microsoft YaHei"
    fr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    document.save(str(output_path))
    return output_path


def word_report_path_for(run_dir: Path) -> Path:
    return run_dir / "report.docx"
